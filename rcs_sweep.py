"""Plot azimuth and/or frequency sweeps from one or more .grim files
and export the plots to a PowerPoint deck."""

import io
import os
import sys

import numpy as np

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

GRIM_DATASET_DIR = '/Users/emery/Documents/GRIM_Revised_2'
if GRIM_DATASET_DIR not in sys.path:
    sys.path.insert(0, GRIM_DATASET_DIR)
from grim_dataset import RcsGrid

# ============================================================================
# USER INPUTS — edit these
# ============================================================================

# Files to overlay (one trace per file on every plot)
FILES = [
    '/Users/emery/Documents/GRIM_Revised_2/rcs_output.grim',
]

# Mode toggles
RUN_AZIMUTH_SWEEP = True
RUN_FREQUENCY_SWEEP = False

# Display unit: None -> first file's default ('dBsm' or 'dBke'), else 'dBsm' or 'dBke'.
DB_UNIT = None

# Output PPTX path
OUTPUT_PPTX = '/Users/emery/Documents/Claude21/rcs_sweep.pptx'

# --- Azimuth sweep (x = azimuth, y = RCS in dB; one subplot per frequency) ---
AZ_XMIN = -90.0
AZ_XMAX = 90.0
AZ_YMIN = -40.0
AZ_YMAX = 30.0
AZ_POLARIZATION = 'VV'           # must match a label in the file (e.g. 'VV','HH','TE','TM')
AZ_FREQUENCIES_GHZ = None        # list, e.g. [2.0, 3.5, 5.0, 7.5, 10.0, 12.0]; None -> first file's frequencies
FREQ_MATCH_TOL_GHZ = 0.01        # tolerance when matching requested freq to file's axis (10 MHz default)

# --- Frequency sweep (x = frequency, y = p50 of RCS in dB across [AZ_MIN, AZ_MAX]) ---
FQ_YMIN = -40.0
FQ_YMAX = 30.0
FQ_FREQUENCIES_GHZ = None        # None -> all frequencies in file; or list e.g. [2.0, 3.0, 5.0]
FQ_AZ_MIN = -30.0
FQ_AZ_MAX = 30.0
FQ_POLARIZATION = 'VV'

# ============================================================================
# Slide layout (10 in x 7.5 in standard 4:3 deck)
SLIDE_W = 10.0
SLIDE_H = 7.5
TITLE_TOP = 0.1
TITLE_H = 0.45
PLOT_AREA_TOP = 0.7
PLOT_AREA_LEFT = 0.30
PLOT_W = 3.10
PLOT_H = 2.40
PLOT_GAP_X = 0.05
PLOT_GAP_Y = 0.05
LEGEND_TOP = PLOT_AREA_TOP + 2 * PLOT_H + PLOT_GAP_Y + 0.10
LEGEND_LEFT = 0.5
LEGEND_W = SLIDE_W - 2 * LEGEND_LEFT
LEGEND_H = SLIDE_H - LEGEND_TOP - 0.2
LABEL_BOX_W = 0.75
LABEL_BOX_H = 0.30
LABEL_INSET = 0.08


def _label_for(path: str) -> str:
    return os.path.splitext(os.path.basename(path))[0]


def _resolve_unit(grid: RcsGrid) -> str:
    return DB_UNIT if DB_UNIT is not None else grid.default_log_unit()


def _linear_to_db(grid: RcsGrid, linear, frequency_value, unit: str) -> np.ndarray:
    if unit.lower() == 'dbke':
        return np.asarray(grid.linear_to_dbke(linear, frequency_value), dtype=float)
    return np.asarray(grid.linear_to_dbsm(linear), dtype=float)


def _check_polarization(grid: RcsGrid, label: str, path: str) -> bool:
    available = [str(p) for p in grid.polarizations]
    if label not in available:
        print(f'[skip] {path}: polarization {label!r} not in {available}')
        return False
    return True


def _match_freq_in_axis(axis_ghz, requested: float, tol: float):
    axis = np.asarray(axis_ghz, dtype=float)
    if axis.size == 0:
        return None
    idx = int(np.argmin(np.abs(axis - float(requested))))
    if abs(axis[idx] - float(requested)) > tol:
        return None
    return float(axis[idx])


def _resolve_az_frequencies():
    if AZ_FREQUENCIES_GHZ is not None:
        return [float(f) for f in AZ_FREQUENCIES_GHZ]
    grid = RcsGrid.load(FILES[0])
    return [float(f) for f in grid.frequencies]


def _format_freq(value: float) -> str:
    return f'{value:g} GHz'


def _fig_to_png_bytes(fig, dpi: int = 200) -> io.BytesIO:
    buf = io.BytesIO()
    fig.savefig(buf, format='png', dpi=dpi, bbox_inches='tight', pad_inches=0.05)
    buf.seek(0)
    return buf


def _build_az_plot(freq_ghz: float, unit: str, files):
    """Render a single azimuth subplot (no title, no legend)."""
    fig, ax = plt.subplots(figsize=(PLOT_W, PLOT_H))
    plotted = False
    for i, path in enumerate(files):
        grid = RcsGrid.load(path)
        if not _check_polarization(grid, AZ_POLARIZATION, path):
            continue
        matched = _match_freq_in_axis(grid.frequencies, freq_ghz, FREQ_MATCH_TOL_GHZ)
        if matched is None:
            print(f'[skip] {path}: {freq_ghz} GHz not on axis '
                  f'{np.asarray(grid.frequencies).tolist()} (tol={FREQ_MATCH_TOL_GHZ})')
            continue
        cropped = grid.axis_crop(
            polarizations=[AZ_POLARIZATION],
            frequencies=[matched],
            azimuth_range=[AZ_XMIN, AZ_XMAX],
        )
        az = np.asarray(cropped.azimuths, dtype=float)
        linear = np.asarray(cropped.rcs_power[:, 0, 0, 0], dtype=float)
        db = _linear_to_db(cropped, linear, matched, unit)
        ax.plot(az, db, color=f'C{i}', linewidth=1.4)
        plotted = True

    ax.set_xlim(AZ_XMIN, AZ_XMAX)
    ax.set_ylim(AZ_YMIN, AZ_YMAX)
    ax.set_xlabel('Azimuth (deg)', fontsize=8)
    ax.set_ylabel(f'RCS ({unit})', fontsize=8)
    ax.tick_params(labelsize=7)
    ax.grid(True, alpha=0.3)
    fig.tight_layout(pad=0.4)
    img = _fig_to_png_bytes(fig)
    plt.close(fig)
    return img, plotted


def _build_legend_image(file_labels):
    fig, ax = plt.subplots(figsize=(8, 0.6))
    ax.axis('off')
    handles = []
    for i, lab in enumerate(file_labels):
        line, = ax.plot([], [], color=f'C{i}', label=lab, linewidth=2.5)
        handles.append(line)
    ax.legend(
        handles=handles,
        loc='center',
        ncol=min(len(file_labels), 4),
        frameon=False,
        fontsize=10,
    )
    img = _fig_to_png_bytes(fig)
    plt.close(fig)
    return img


def _slide_position(row: int, col: int):
    x = PLOT_AREA_LEFT + col * (PLOT_W + PLOT_GAP_X)
    y = PLOT_AREA_TOP + row * (PLOT_H + PLOT_GAP_Y)
    return x, y


def _add_freq_label_box(slide, plot_x: float, plot_y: float, text: str) -> None:
    x = plot_x + LABEL_INSET
    y = plot_y + LABEL_INSET
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y),
        Inches(LABEL_BOX_W), Inches(LABEL_BOX_H),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0xFF, 0x00, 0x00)
    shape.line.color.rgb = RGBColor(0xFF, 0x00, 0x00)
    tf = shape.text_frame
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.0)
    tf.margin_bottom = Inches(0.0)
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = text
    run = p.runs[0]
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    run.font.bold = True
    run.font.size = Pt(11)


def _new_blank_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])  # 'Blank' layout


def _add_title(slide, text: str) -> None:
    box = slide.shapes.add_textbox(
        Inches(0.3), Inches(TITLE_TOP),
        Inches(SLIDE_W - 0.6), Inches(TITLE_H),
    )
    tf = box.text_frame
    tf.text = text
    run = tf.paragraphs[0].runs[0]
    run.font.size = Pt(18)
    run.font.bold = True


def azimuth_sweep_to_pptx(prs: Presentation) -> int:
    requested = _resolve_az_frequencies()
    if not requested:
        print('Azimuth sweep: no frequencies requested.')
        return 0

    sample = RcsGrid.load(FILES[0])
    unit = _resolve_unit(sample)

    file_units = {p: _resolve_unit(RcsGrid.load(p)) for p in FILES}
    if len(set(file_units.values())) > 1:
        print(f'[warn] mixed default units across files: {file_units}; using {unit}')

    rendered = []
    for f in requested:
        img, did = _build_az_plot(f, unit, FILES)
        if did:
            rendered.append((f, img))
        else:
            print(f'[skip] no traces at {f} GHz')

    if not rendered:
        print('Azimuth sweep: no plots generated.')
        return 0

    legend_img = _build_legend_image([_label_for(p) for p in FILES])

    slides_added = 0
    for start in range(0, len(rendered), 6):
        chunk = rendered[start:start + 6]
        slide = _new_blank_slide(prs)
        _add_title(slide, f'Azimuth sweep — pol={AZ_POLARIZATION} ({unit})')
        for k, (freq, img) in enumerate(chunk):
            row, col = divmod(k, 3)
            x, y = _slide_position(row, col)
            img.seek(0)
            slide.shapes.add_picture(
                img, Inches(x), Inches(y),
                Inches(PLOT_W), Inches(PLOT_H),
            )
            _add_freq_label_box(slide, x, y, _format_freq(freq))
        legend_img.seek(0)
        slide.shapes.add_picture(
            legend_img, Inches(LEGEND_LEFT), Inches(LEGEND_TOP),
            Inches(LEGEND_W), Inches(LEGEND_H),
        )
        slides_added += 1
    return slides_added


def frequency_sweep_to_pptx(prs: Presentation) -> int:
    sample = RcsGrid.load(FILES[0])
    unit = _resolve_unit(sample)

    fig, ax = plt.subplots(figsize=(9, 5))
    plotted = 0
    for i, path in enumerate(FILES):
        grid = RcsGrid.load(path)
        if not _check_polarization(grid, FQ_POLARIZATION, path):
            continue
        available = np.asarray(grid.frequencies, dtype=float)
        if FQ_FREQUENCIES_GHZ is None:
            freqs = available
        else:
            picked = []
            for f in FQ_FREQUENCIES_GHZ:
                m = _match_freq_in_axis(available, float(f), FREQ_MATCH_TOL_GHZ)
                if m is not None:
                    picked.append(m)
            freqs = np.asarray(sorted(set(picked)), dtype=float)
            if freqs.size == 0:
                print(f'[skip] {path}: requested frequencies not on axis {available.tolist()}.')
                continue
        cropped = grid.axis_crop(
            polarizations=[FQ_POLARIZATION],
            frequencies=list(freqs.tolist()),
            azimuth_range=[FQ_AZ_MIN, FQ_AZ_MAX],
        )
        f_axis = np.asarray(cropped.frequencies, dtype=float)
        p50 = np.empty(f_axis.shape, dtype=float)
        for j, f_val in enumerate(f_axis):
            linear = np.asarray(cropped.rcs_power[:, 0, j, 0], dtype=float)
            db = _linear_to_db(cropped, linear, f_val, unit)
            p50[j] = float(np.nanpercentile(db, 50.0))
        ax.plot(f_axis, p50, marker='o', color=f'C{i}', label=_label_for(path))
        plotted += 1

    if plotted == 0:
        plt.close(fig)
        print('Frequency sweep: nothing to plot.')
        return 0

    ax.set_ylim(FQ_YMIN, FQ_YMAX)
    ax.set_xlabel('Frequency (GHz)')
    ax.set_ylabel(f'p50 RCS ({unit}) over az [{FQ_AZ_MIN}°, {FQ_AZ_MAX}°]')
    ax.set_title(f'Frequency sweep — pol={FQ_POLARIZATION}')
    ax.grid(True, alpha=0.3)
    ax.legend(loc='best', fontsize=9)
    fig.tight_layout()
    img = _fig_to_png_bytes(fig)
    plt.close(fig)

    slide = _new_blank_slide(prs)
    _add_title(slide, f'Frequency sweep — pol={FQ_POLARIZATION} ({unit})')
    slide.shapes.add_picture(
        img, Inches(0.5), Inches(0.7),
        Inches(SLIDE_W - 1.0), Inches(SLIDE_H - 1.0),
    )
    return 1


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    added = 0
    if RUN_AZIMUTH_SWEEP:
        added += azimuth_sweep_to_pptx(prs)
    if RUN_FREQUENCY_SWEEP:
        added += frequency_sweep_to_pptx(prs)

    if added == 0:
        print('No slides generated; PPTX not written.')
        return

    prs.save(OUTPUT_PPTX)
    print(f'Wrote {OUTPUT_PPTX} ({added} slide(s)).')


if __name__ == '__main__':
    main()
